from __future__ import annotations

import shutil
import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from eismaster.analysis.fitting import FitOutcome
from eismaster.analysis.quality import assess_spectrum_quality
from eismaster.exporters import (
    build_drt_matrices,
    export_batch_summary,
    export_fit_results,
    export_spectrum_bundle,
    write_drt_only_export,
)
from eismaster.io.chi import load_spectrum
from eismaster.models import BatchItemResult, BatchSummary, SpectrumData, SpectrumMetadata
from tests.fixture_factory import write_single_arc_txt

ROOT = Path(__file__).resolve().parents[1]


def _dummy_spectrum(stem: str = "Ag_EIS_OCV") -> SpectrumData:
    return SpectrumData(
        metadata=SpectrumMetadata(file_path=Path(f"{stem}.txt")),
        freq_hz=np.array([100.0, 10.0, 1.0]),
        z_real_ohm=np.array([1.0, 2.0, 3.0]),
        z_imag_ohm=np.array([-0.1, -0.2, -0.3]),
        z_mod_ohm=np.array([1.0, 2.0, 3.0]),
        phase_deg=np.array([0.0, 0.0, 0.0]),
    )


class ExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp_root = ROOT / "tests" / "_tmp_exports"
        if self.tmp_root.exists():
            shutil.rmtree(self.tmp_root)
        self.tmp_root.mkdir(parents=True, exist_ok=True)
        self.txt_sample = write_single_arc_txt(self.tmp_root / "Ag_EIS_OCV.txt")

    def tearDown(self) -> None:
        if self.tmp_root.exists():
            shutil.rmtree(self.tmp_root)

    def test_txt_bundle_has_expected_headers(self) -> None:
        spectrum = load_spectrum(self.txt_sample)
        paths = export_spectrum_bundle(self.tmp_root / "bundle", spectrum)
        plot_text = paths["raw_plot"].read_text(encoding="utf-8").splitlines()[0]
        rs_rct_text = paths["rs_rct"].read_text(encoding="utf-8").splitlines()[0]
        fit_text = paths["fit_overlay"].read_text(encoding="utf-8").splitlines()[0]
        fit_second = paths["fit_overlay"].read_text(encoding="utf-8").splitlines()[1]
        self.assertEqual(plot_text, "OCV\t")
        self.assertEqual(rs_rct_text, "label\tfile\tRs\tRct\tRsei")
        self.assertEqual(fit_text, "OCV\t\t\t")
        self.assertEqual(fit_second, "z_real_exp\timag_exp_pos\tz_real_fit\timag_fit_pos")

    def test_fit_and_batch_exports(self) -> None:
        spectrum = load_spectrum(self.txt_sample)
        quality = assess_spectrum_quality(spectrum)
        fit = FitOutcome(
            model_key="demo",
            model_label="Demo Model",
            status="ok",
            message="ok",
            parameters={"Rs": 5.8, "Rct": 120.0},
            statistics={"rss": 1.2, "aic": 2.3, "bic": 3.4, "chi2_reduced": 0.12},
            predicted_real_ohm=spectrum.z_real_ohm.copy(),
            predicted_imag_ohm=spectrum.z_imag_ohm.copy(),
        )
        summary = BatchSummary(model_key="demo", items=[BatchItemResult(spectrum=spectrum, quality=quality, fit=fit)])
        fit_paths = export_fit_results(self.tmp_root / "fit", spectrum, quality, fit)
        summary_paths = export_batch_summary(self.tmp_root / "fit", summary)
        params_header = fit_paths["rs_rct"].read_text(encoding="utf-8").splitlines()[0]
        batch_header = summary_paths["rs_rct"].read_text(encoding="utf-8").splitlines()[0]
        curve_header = fit_paths["fit_overlay"].read_text(encoding="utf-8").splitlines()[0]
        self.assertIn("label\tfile\tRs\tRct", params_header)
        self.assertIn("label\tfile\tRs\tRct", batch_header)
        self.assertEqual(curve_header, "OCV\t\t\t")

    def test_xlsx_export_creates_workbook(self) -> None:
        spectrum = load_spectrum(self.txt_sample)
        quality = assess_spectrum_quality(spectrum)
        fit = FitOutcome(
            model_key="demo",
            model_label="Demo Model",
            status="ok",
            message="ok",
            parameters={"Rs": 5.8, "Rct": 120.0},
            statistics={"rss": 1.2, "aic": 2.3, "bic": 3.4, "chi2_reduced": 0.12},
            predicted_real_ohm=spectrum.z_real_ohm.copy(),
            predicted_imag_ohm=spectrum.z_imag_ohm.copy(),
        )
        xlsx_path = self.tmp_root / "export.xlsx"
        paths = export_spectrum_bundle(xlsx_path, spectrum, fit=fit, quality=quality, fmt="xlsx")
        self.assertIn("workbook", paths)
        self.assertTrue(paths["workbook"].exists())
        self.assertEqual(paths["workbook"].suffix, ".xlsx")

    def test_xlsx_batch_export(self) -> None:
        spectrum = load_spectrum(self.txt_sample)
        quality = assess_spectrum_quality(spectrum)
        fit = FitOutcome(
            model_key="demo",
            model_label="Demo Model",
            status="ok",
            message="ok",
            parameters={"Rs": 5.8, "Rct": 120.0},
            statistics={"rss": 1.2, "aic": 2.3, "bic": 3.4, "chi2_reduced": 0.12},
            predicted_real_ohm=spectrum.z_real_ohm.copy(),
            predicted_imag_ohm=spectrum.z_imag_ohm.copy(),
        )
        summary = BatchSummary(model_key="demo", items=[BatchItemResult(spectrum=spectrum, quality=quality, fit=fit)])
        xlsx_path = self.tmp_root / "batch.xlsx"
        paths = export_batch_summary(xlsx_path, summary, fmt="xlsx")
        self.assertIn("workbook", paths)
        self.assertTrue(paths["workbook"].exists())

    def test_drt_export_builds_line_density_and_area_matrices(self) -> None:
        spectrum = _dummy_spectrum()
        drt_path = self.tmp_root / f"{spectrum.metadata.file_path.stem}_DRT.txt"
        drt_path.write_text("tau\tgamma(tau)\n1\t1\n10\t2\n100\t3\n", encoding="utf-8")

        matrices = build_drt_matrices([spectrum], self.tmp_root)

        self.assertEqual(set(matrices), {"drt_line_plot", "drt_cloud_density", "drt_quant_area"})
        self.assertEqual(matrices["drt_line_plot"].iloc[1].tolist(), ["logtau", "gamma_tau"])
        self.assertEqual(matrices["drt_line_plot"].iloc[1, 1], "gamma_tau")
        self.assertEqual(matrices["drt_cloud_density"].iloc[1, 1], "gamma_tau")
        self.assertAlmostEqual(float(matrices["drt_cloud_density"].iloc[3, 1]), 2.0)
        self.assertEqual(
            matrices["drt_quant_area"].columns.tolist(),
            [
                "sample",
                "time_min",
                "total_area",
                "logtau_lt_neg3",
                "logtau_neg3_to_0",
                "logtau_ge_0",
            ],
        )
        self.assertAlmostEqual(float(matrices["drt_quant_area"].iloc[0]["total_area"]), 13.815510557964274)
        self.assertAlmostEqual(float(matrices["drt_quant_area"].iloc[0]["logtau_ge_0"]), 13.815510557964274)

    def test_drt_quant_area_accepts_custom_logtau_breaks(self) -> None:
        spectrum = _dummy_spectrum("Ag_S01_EIS_T12M")
        drt_path = self.tmp_root / f"{spectrum.metadata.file_path.stem}_DRT.txt"
        drt_path.write_text("tau\tgamma(tau)\n1\t1\n10\t2\n100\t3\n", encoding="utf-8")

        matrices = build_drt_matrices([spectrum], self.tmp_root, logtau_breaks=(1.0, 3.0))
        trend = matrices["drt_quant_area"]

        self.assertEqual(
            trend.columns.tolist(),
            ["sample", "time_min", "total_area", "logtau_lt_1", "logtau_1_to_3", "logtau_ge_3"],
        )
        self.assertEqual(float(trend.iloc[0]["time_min"]), 12.0)
        self.assertAlmostEqual(float(trend.iloc[0]["logtau_lt_1"]), 2.302585092994046)
        self.assertAlmostEqual(float(trend.iloc[0]["logtau_1_to_3"]), 11.512925464970229)
        self.assertAlmostEqual(float(trend.iloc[0]["logtau_ge_3"]), 0.0)

    def test_drt_line_plot_can_export_tau_seconds_axis(self) -> None:
        spectrum = _dummy_spectrum()
        drt_path = self.tmp_root / f"{spectrum.metadata.file_path.stem}_DRT.txt"
        drt_path.write_text("tau\tgamma(tau)\n1\t1\n10\t2\n100\t3\n", encoding="utf-8")

        matrices = build_drt_matrices([spectrum], self.tmp_root, line_x_axis="tau")
        line_plot = matrices["drt_line_plot"]

        self.assertEqual(line_plot.iloc[1].tolist(), ["tau/s", "gamma_tau"])
        self.assertEqual(float(line_plot.iloc[2, 0]), 1.0)
        self.assertEqual(float(line_plot.iloc[3, 0]), 10.0)
        self.assertEqual(float(line_plot.iloc[4, 0]), 100.0)
        self.assertEqual(float(line_plot.iloc[3, 1]), 2.0)

    def test_drt_matrix_accepts_peak_fit_exports(self) -> None:
        spectrum = _dummy_spectrum()
        drt_path = self.tmp_root / f"{spectrum.metadata.file_path.stem}_DRT.txt"
        drt_path.write_text(
            "\n".join(
                [
                    "L\t0",
                    "R\t1",
                    "function\tpeak_height*exp(-1/2*(log_tau - peak_position)^2/peak_width^2)",
                    "peak number\tpeak height\tpeak position\tpeak width",
                    "1\t2\t0\t0.5",
                ]
            ),
            encoding="utf-8",
        )

        matrices = build_drt_matrices([spectrum], self.tmp_root)

        self.assertIn("drt_quant_area", matrices)
        peak_area = float(matrices["drt_quant_area"].iloc[0]["total_area"])
        self.assertAlmostEqual(peak_area, 2.5066, places=2)

    def test_drt_only_export_writes_one_workbook_with_three_sheets(self) -> None:
        spectrum = _dummy_spectrum()
        drt_path = self.tmp_root / f"{spectrum.metadata.file_path.stem}_DRT.txt"
        drt_path.write_text("tau\tgamma(tau)\n1\t1\n10\t2\n100\t3\n", encoding="utf-8")

        paths = write_drt_only_export(self.tmp_root / "drt_matrix.xlsx", [spectrum], self.tmp_root)

        self.assertEqual(set(paths), {"drt"})
        self.assertTrue(paths["drt"].exists())
        workbook = load_workbook(paths["drt"], read_only=True)
        self.assertEqual(set(workbook.sheetnames), {"drt_line_plot", "drt_cloud_density", "drt_quant_area"})
        quant_sheet = workbook["drt_quant_area"]
        self.assertEqual(quant_sheet["A1"].value, "sample")
        self.assertEqual(quant_sheet["C1"].value, "total_area")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
